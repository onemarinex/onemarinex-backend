"""Fleet register, effective-dated declarations and private, resumable imports."""

from __future__ import annotations
from typing import Optional
import csv
import hashlib
import io
from datetime import datetime, date, timezone
from pathlib import PurePath
from zipfile import ZipFile
from typing import Literal
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from fastapi.responses import Response
from pydantic import BaseModel, Field, ValidationError
from sqlalchemy import func, or_, and_
from sqlalchemy.orm import Session
from app.db.session import get_db
from app.db.models.vessel import Vessel
from app.db.models.shipping_company import VesselRoleClaim
from app.db.models.company_fleet_import import CompanyFleetImport
from app.db.models.company_fleet_change import CompanyFleetChange
from app.api.v1.routes_auth import get_current_user
from app.api.v1.routes_shipping_company import (
    current_company_member,
    _require_company_admin,
)
from app.services import company_fleet_declarations as fleet
from app.services.company_claims import _interval

router = APIRouter()
LABELS = [
    "IMO",
    "Vessel Name",
    "Owner",
    "Commercial Manager",
    "ISM Manager",
    "Start Date",
    "End Date",
]


def write_context(context):
    _require_company_admin(context[1])
    return context[0]


@router.get("/resolve/{imo_number}")
def resolve(
    imo_number: str,
    db: Session = Depends(get_db),
    context=Depends(current_company_member),
):
    try:
        number = fleet.imo(imo_number)
        vessel = fleet.resolve(db, number)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from None
    return {"imo_number": number, "vessel_name": vessel.name}


@router.get("/entries")
def entries(
    search: str = Query("", max_length=120),
    view: Literal["active", "historical", "all"] = "active",
    role: Optional[
        Literal["REGISTERED_OWNER", "COMMERCIAL_MANAGER", "ISM_MANAGER"]
    ] = None,
    sort: Literal["name", "imo", "start", "end"] = "name",
    descending: bool = False,
    offset: int = Query(0, ge=0),
    limit: int = Query(25, ge=1, le=100),
    db: Session = Depends(get_db),
    context=Depends(current_company_member),
):
    q, group = fleet.group_query(db, context[0].id)
    now = datetime.now(timezone.utc)
    live = and_(
        VesselRoleClaim.status.in_(fleet.ACCEPTED),
        VesselRoleClaim.revoked_at.is_(None),
        VesselRoleClaim.effective_start <= now,
        or_(
            VesselRoleClaim.effective_end.is_(None), VesselRoleClaim.effective_end > now
        ),
    )
    current_count = (
        q.filter(live)
        .with_entities(func.count(func.distinct(VesselRoleClaim.vessel_id)))
        .scalar()
    )
    q = q.join(Vessel, Vessel.id == VesselRoleClaim.vessel_id)
    if search.strip():
        term = "%" + search.strip().replace("%", "\\%").replace("_", "\\_") + "%"
        q = q.filter(
            or_(
                Vessel.name.ilike(term, escape="\\"),
                Vessel.imo_number.ilike(term, escape="\\"),
                VesselRoleClaim.declared_vessel_name.ilike(term, escape="\\"),
            )
        )
    groups = q.with_entities(
        group.label("entry_group"), func.min(VesselRoleClaim.id).label("id")
    ).group_by(group)
    if view != "all":
        groups = groups.having(
            func.bool_or(live) if view == "active" else ~func.bool_or(live)
        )
    if role:
        groups = groups.having(func.bool_or(VesselRoleClaim.role == role))
    order = {
        "name": func.min(
            func.coalesce(VesselRoleClaim.declared_vessel_name, Vessel.name)
        ),
        "imo": func.min(Vessel.imo_number),
        "start": func.min(VesselRoleClaim.effective_start),
        "end": func.max(VesselRoleClaim.effective_end),
    }[sort]
    total = groups.count()
    keys = (
        groups.order_by(
            order.desc().nullslast() if descending else order.asc().nullslast(),
            func.min(VesselRoleClaim.id),
        )
        .offset(offset)
        .limit(limit)
        .all()
    )
    return dict(
        total=total,
        offset=offset,
        limit=limit,
        current_vessels=current_count,
        items=[
            fleet.entry_out(db, fleet.load_group(db, context[0].id, k.id)) for k in keys
        ],
    )


@router.post("/entries", status_code=201)
def add_entry(
    body: fleet.DeclarationIn,
    db: Session = Depends(get_db),
    context=Depends(current_company_member),
    actor=Depends(get_current_user),
):
    company = write_context(context)
    try:
        rows = fleet.declare(db, company, actor, body)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from None
    db.commit()
    return fleet.entry_out(db, rows)


class EditIn(BaseModel):
    model_config = {"str_strip_whitespace": True, "extra": "forbid"}
    action: Literal["ROLES", "END", "VOID", "TYPO", "RENAME", "IMO"]
    revision: str = Field(min_length=64, max_length=64)
    effective_date: Optional[str] = Field(default=None, max_length=40)
    end_date: Optional[str] = Field(default=None, max_length=40)
    owner: bool = False
    commercial_manager: bool = False
    ism_manager: bool = False
    corrected_name: Optional[str] = Field(default=None, min_length=1, max_length=255)
    corrected_imo: Optional[str] = Field(default=None, max_length=32)
    reason: str = Field(min_length=10, max_length=4000)


def apply_change(db, company, actor, rows, body):
    if any(c.status not in fleet.OPEN or c.revoked_at for c in rows):
        raise HTTPException(409, "This entry is closed. Add a new period if needed.")
    first = rows[0]
    try:
        point = fleet.boundary(body.effective_date)
        end = fleet.boundary(body.end_date)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from None
    if body.action in ("END", "ROLES") and (
        point is None
        or point <= first.effective_start
        or (first.effective_end and point >= first.effective_end)
    ):
        raise HTTPException(
            422, "Choose an effective date inside the current period, after its start."
        )
    if body.action == "ROLES" and not any(
        getattr(body, field) for field in fleet.ROLE_FIELDS
    ):
        raise HTTPException(422, "Select at least one company role.")
    if body.action == "TYPO" and (point is None or (end and end <= point)):
        raise HTTPException(422, "Enter a valid corrected period.")
    if body.action == "RENAME" and not (body.corrected_name or "").strip():
        raise HTTPException(422, "Enter the corrected vessel name.")
    if body.action == "IMO":
        try:
            target = fleet.resolve(db, fleet.imo(body.corrected_imo))
        except ValueError as exc:
            raise HTTPException(422, str(exc)) from None
        if target.id == first.vessel_id:
            raise HTTPException(422, "Select a different canonical IMO.")
    # Verified periods, identity corrections and administrative typos need independent review.
    needs_review = any(c.status == "APPROVED" for c in rows) or body.action in (
        "TYPO",
        "RENAME",
        "IMO",
    )
    if needs_review:
        if (
            db.query(CompanyFleetChange.id)
            .filter(
                CompanyFleetChange.claim_id.in_([c.id for c in rows]),
                CompanyFleetChange.status == "REQUESTED",
            )
            .first()
        ):
            raise HTTPException(409, "A change for this entry already awaits review.")
        change = CompanyFleetChange(
            company_id=company.id,
            claim_id=first.id,
            requested_by_user_id=actor.id,
            action="GROUP",
            payload=body.model_dump(),
            reason=body.reason.strip(),
        )
        db.add(change)
        db.flush()
        fleet.audit(
            db,
            actor,
            company.id,
            first.id,
            "fleet_group_change_requested",
            after=body.model_dump(),
            reason=body.reason,
        )
        return {"status": "REQUESTED", "change_id": change.id}
    return apply_direct(db, company, actor, rows, body)


def apply_direct(db, company, actor, rows, body):
    first = rows[0]
    old_end = first.effective_end
    before = [_interval(c) for c in rows]
    point = fleet.boundary(body.effective_date)
    # Only declaration records can be changed without SA approval.
    if body.action == "VOID":
        for c in rows:
            c.status = "REVOKED"
            c.revoked_at = datetime.now(timezone.utc)
    elif body.action in ("END", "ROLES"):
        for c in rows:
            c.effective_end = point
    else:
        raise HTTPException(422, "This change requires Super Admin review.")
    db.flush()
    result = None
    if body.action == "ROLES":
        row = fleet.DeclarationIn(
            imo_number=first.imo_number,
            vessel_name=db.get(Vessel, first.vessel_id).name,
            **{f: getattr(body, f) for f in fleet.ROLE_FIELDS},
            start_date=point.isoformat(),
            end_date=old_end.isoformat() if old_end else None,
        )
        try:
            result = fleet.declare(db, company, actor, row)
        except ValueError as exc:
            raise HTTPException(409, str(exc)) from None
    fleet.audit(
        db,
        actor,
        company.id,
        first.id,
        "fleet_" + body.action.lower(),
        before={"claims": before},
        after={"claims": [_interval(c) for c in rows]},
        reason=body.reason.strip(),
    )
    return {"status": "APPLIED", "entry": fleet.entry_out(db, result or rows)}


def apply_reviewed_change(db, change, actor, decision):
    from app.db.models.shipping_company import ShippingCompany

    first = db.get(VesselRoleClaim, change.claim_id)
    body = EditIn(**change.payload)
    # Lock both old and corrected canonical identities in a stable order.
    target = None
    if body.action == "IMO":
        try:
            target = fleet.resolve(db, fleet.imo(body.corrected_imo))
        except ValueError as exc:
            raise HTTPException(422, str(exc)) from None
    ids = [first.vessel_id] + ([target.id] if target else [])
    db.query(Vessel).filter(Vessel.id.in_(ids)).order_by(
        Vessel.id
    ).with_for_update().all()
    rows = fleet.load_group(db, change.company_id, first.id, lock=True)
    if fleet.revision(rows) != body.revision:
        raise HTTPException(
            409, "The entry changed after this request. Ask for a new request."
        )
    if not (decision.evidence_reference or "").strip():
        raise HTTPException(422, "Record the independent evidence for this change.")
    company = db.get(ShippingCompany, change.company_id)
    before = {"claims": [_interval(c) for c in rows]}
    if body.action == "RENAME":
        vessel = db.get(Vessel, first.vessel_id)
        old_name = vessel.name
        vessel.name = body.corrected_name.strip()
        fleet.audit(
            db,
            actor,
            company.id,
            first.id,
            "vessel_master_name_corrected",
            before={"name": old_name},
            after={"name": vessel.name},
            reason=decision.reason + " Evidence: " + decision.evidence_reference,
        )
        return
    if body.action in ("END", "VOID", "ROLES"):
        # SA approval permits the same history-preserving operation on verified entries.
        result = apply_direct(db, company, actor, rows, body)
        if body.action == "ROLES":
            replacements = fleet.load_group(db, company.id, result["entry"]["id"])
            for c in replacements:
                if c.status == "AUTO_ACCEPTED":
                    c.status = "APPROVED"
                    c.approved_by_user_id = actor.id
                    c.approved_at = datetime.now(timezone.utc)
                    c.evidence_type = "independent_fleet_review"
                    c.evidence_reference = decision.evidence_reference
        fleet.audit(
            db,
            actor,
            company.id,
            first.id,
            "reviewed_fleet_change_applied",
            before=before,
            reason=decision.reason + " Evidence: " + decision.evidence_reference,
        )
        return
    # A typo or wrong IMO never moves a port call, booking or assignment.
    for c in rows:
        c.status = "CORRECTED"
        c.revoked_at = datetime.now(timezone.utc)
    db.flush()
    vessel = target or db.get(Vessel, first.vessel_id)
    start = (
        fleet.boundary(body.effective_date)
        if body.action == "TYPO"
        else first.effective_start
    )
    end = (
        fleet.boundary(body.end_date) if body.action == "TYPO" else first.effective_end
    )
    row = fleet.DeclarationIn(
        imo_number=vessel.imo_number,
        vessel_name=vessel.name,
        **{
            f: any(c.role == code for c in rows)
            for f, code in fleet.ROLE_FIELDS.items()
        },
        start_date=start.isoformat(),
        end_date=end.isoformat() if end else None,
    )
    try:
        replacements = fleet.declare(db, company, actor, row)
    except ValueError as exc:
        raise HTTPException(409, str(exc)) from None
    change.replacement_claim_id = replacements[0].id
    fleet.audit(
        db,
        actor,
        company.id,
        first.id,
        "reviewed_fleet_identity_or_date_corrected",
        before=before,
        after={"claims": [_interval(c) for c in replacements]},
        reason=decision.reason + " Evidence: " + decision.evidence_reference,
    )


@router.post("/entries/{entry_id}/change")
def change_entry(
    entry_id: int,
    body: EditIn,
    db: Session = Depends(get_db),
    context=Depends(current_company_member),
    actor=Depends(get_current_user),
):
    company = write_context(context)
    # Canonical row lock always precedes claim locks, including bulk imports.
    first = (
        db.query(VesselRoleClaim).filter_by(id=entry_id, company_id=company.id).first()
    )
    if first is None:
        raise HTTPException(404, "Fleet entry not found.")
    db.query(Vessel).filter_by(id=first.vessel_id).with_for_update().one()
    rows = fleet.load_group(db, company.id, entry_id, lock=True)
    if fleet.revision(rows) != body.revision:
        raise HTTPException(409, "The fleet entry changed. Refresh and try again.")
    result = apply_change(db, company, actor, rows, body)
    db.commit()
    return result


def safe_csv(values):
    output = io.StringIO()
    writer = csv.writer(output)
    for row in values:
        writer.writerow(
            [
                "'" + str(v) if str(v).lstrip().startswith(("=", "+", "-", "@")) else v
                for v in row
            ]
        )
    return output.getvalue()


@router.get("/template.csv")
def template(context=Depends(current_company_member)):
    return Response(
        safe_csv([LABELS]),
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="fleet-template.csv"'},
    )


def draft_out(draft):
    return dict(
        id=draft.id,
        filename=draft.filename,
        source_sha256=draft.source_sha256,
        columns=draft.columns,
        rows=draft.rows,
        mapping=draft.mapping,
        status=draft.status,
        result=draft.result,
        imported_at=draft.imported_at,
    )


def owned_import(db, company_id, import_id, lock=False):
    q = db.query(CompanyFleetImport).filter_by(id=import_id, company_id=company_id)
    row = (q.with_for_update() if lock else q).first()
    if row is None:
        raise HTTPException(404, "Fleet upload not found.")
    return row


@router.post("/imports", status_code=201)
def upload(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    context=Depends(current_company_member),
    actor=Depends(get_current_user),
):
    company = write_context(context)
    raw = file.file.read(2_000_001)
    if len(raw) > 2_000_000:
        raise HTTPException(413, "Use a file of at most 2 MB.")
    filename = PurePath((file.filename or "fleet.csv").replace("\\", "/")).name[:255]
    try:
        if filename.lower().endswith(".csv"):
            reader = csv.reader(io.StringIO(raw.decode("utf-8-sig")))
            data = []
            for r in reader:
                if len(r) > 32:
                    raise ValueError("Use at most 32 columns.")
                if any(str(v).strip() for v in r):
                    data.append(r)
                if len(data) > 501:
                    raise ValueError("Use at most 500 data rows.")
        elif filename.lower().endswith(".xlsx"):
            from openpyxl import load_workbook

            with ZipFile(io.BytesIO(raw)) as archive:
                if (
                    len(archive.infolist()) > 1000
                    or sum(x.file_size for x in archive.infolist()) > 8_000_000
                ):
                    raise ValueError("The workbook expands beyond 8 MB.")
            book = load_workbook(
                io.BytesIO(raw), read_only=True, data_only=False, keep_links=False
            )
            try:
                sheet = book.active
                if (sheet.max_row or 0) > 501 or (sheet.max_column or 0) > 32:
                    raise ValueError("Use at most 500 rows and 32 columns.")
                data = [
                    [
                        (
                            v.isoformat()
                            if isinstance(v, (date, datetime))
                            else str(v) if v is not None else ""
                        )
                        for v in row
                    ]
                    for row in sheet.iter_rows(
                        max_row=501,
                        max_col=min(sheet.max_column or 1, 32),
                        values_only=True,
                    )
                ]
                data = [r for r in data if any(v.strip() for v in r)]
            finally:
                book.close()
        else:
            raise ValueError("Use a CSV or XLSX file.")
        if len(data) < 2:
            raise ValueError("Include a header and at least one data row.")
        if any(len(str(v)) > 4000 for row in data for v in row):
            raise ValueError("A cell exceeds 4,000 characters.")
    except (ValueError, UnicodeError) as exc:
        raise HTTPException(422, str(exc)) from None
    except Exception:
        raise HTTPException(422, "The file could not be read.") from None
    columns = [str(v).strip() or f"Column {i+1}" for i, v in enumerate(data[0])]
    mapping = {
        field: columns.index(label)
        for field, label in zip(fleet.FIELDS, LABELS)
        if label in columns
    }
    draft = CompanyFleetImport(
        company_id=company.id,
        created_by_user_id=actor.id,
        filename=filename,
        source_sha256=hashlib.sha256(raw).hexdigest(),
        columns=columns,
        rows=data[1:],
        mapping=mapping,
    )
    db.add(draft)
    db.flush()
    fleet.audit(
        db,
        actor,
        company.id,
        draft.id,
        "fleet_file_uploaded",
        target_type="fleet_import",
        after={
            "filename": filename,
            "sha256": draft.source_sha256,
            "rows": len(draft.rows),
        },
    )
    db.commit()
    return draft_out(draft)


@router.get("/imports/{import_id}")
def read_import(
    import_id: int,
    db: Session = Depends(get_db),
    context=Depends(current_company_member),
):
    company = write_context(context)
    return draft_out(owned_import(db, company.id, import_id))


class MappingIn(BaseModel):
    mapping: dict[str, int]


def parse_rows(db, company_id, draft, mapping):
    if (
        set(mapping) != set(fleet.FIELDS)
        or len(set(mapping.values())) != 7
        or any(i < 0 or i >= len(draft.columns) for i in mapping.values())
    ):
        raise HTTPException(
            422, "Map each of the seven fields to a different source column."
        )
    findings = []
    parsed = []
    seen = []
    for index, cells in enumerate(draft.rows, 2):
        values = {
            field: str(cells[i] if i < len(cells) else "").strip()
            for field, i in mapping.items()
        }
        state = "INVALID"
        message = ""
        row = None
        try:
            if any(v.startswith(("=", "+", "@")) for v in values.values()):
                raise ValueError("Formula cells are not allowed.")
            for field in fleet.ROLE_FIELDS:
                val = values[field].casefold()
                if val not in ("yes", "no", "true", "false", "1", "0"):
                    raise ValueError("Role fields must contain Yes or No.")
                values[field] = val in ("yes", "true", "1")
            # Excel dates have no zone: interpret calendar midnight as the agreed company boundary.
            for field in ("start_date", "end_date"):
                val = values[field]
                if isinstance(val, str) and val.endswith("T00:00:00"):
                    values[field] = val[:10]
            row = fleet.DeclarationIn(**values)
            vessel, state, _ = fleet.inspect_row(db, company_id, row)
            start, end = fleet.boundary(row.start_date), fleet.boundary(row.end_date)
            if any(
                number == row.imo_number
                and set(roles) & set(row.roles())
                and (other_end is None or start < other_end)
                and (end is None or other_start < end)
                for number, roles, other_start, other_end in seen
            ):
                state = "DUPLICATE"
            seen.append((row.imo_number, row.roles(), start, end))
            message = {
                "ACCEPTED": "Ready for aggregate acceptance.",
                "CONFLICT": "Pending Super Admin review. Other valid rows can proceed.",
                "DUPLICATE": "An overlapping role already exists in your fleet or this file.",
            }[state]
        except (ValueError, ValidationError) as exc:
            text = str(exc)
            state = "UNKNOWN" if "Unknown or ambiguous IMO" in text else "INVALID"
            message = (
                text.split("\n")[0]
                if not isinstance(exc, ValidationError)
                else "Check IMO, roles, vessel name and effective dates."
            )
        findings.append(
            dict(
                row=index,
                imo_number=str(values.get("imo_number", "")),
                status=state,
                message=message,
            )
        )
        parsed.append(row if state in ("ACCEPTED", "CONFLICT") else None)
    return findings, parsed


def preview_result(draft, findings):
    return dict(
        import_id=draft.id,
        total=len(findings),
        counts={
            state.lower(): sum(f["status"] == state for f in findings)
            for state in ("ACCEPTED", "INVALID", "DUPLICATE", "UNKNOWN", "CONFLICT")
        },
        findings=findings,
    )


@router.post("/imports/{import_id}/preview")
def preview(
    import_id: int,
    body: MappingIn,
    db: Session = Depends(get_db),
    context=Depends(current_company_member),
):
    company = write_context(context)
    draft = owned_import(db, company.id, import_id, True)
    if draft.status == "IMPORTED":
        return draft.result
    findings, _ = parse_rows(db, company.id, draft, body.mapping)
    draft.mapping = body.mapping
    draft.status = "REVIEWED"
    db.commit()
    return preview_result(draft, findings)


@router.post("/imports/{import_id}/confirm")
def confirm(
    import_id: int,
    db: Session = Depends(get_db),
    context=Depends(current_company_member),
    actor=Depends(get_current_user),
):
    company = write_context(context)
    draft = owned_import(db, company.id, import_id, True)
    if draft.status == "IMPORTED":
        return draft.result
    if draft.status != "REVIEWED":
        raise HTTPException(409, "Review the mapped rows before import.")
    fleet.require_fleet_constraint(db)
    # All canonical vessel locks in one stable order prevent deadlocks between bulk files.
    numbers = []
    for cells in draft.rows:
        try:
            numbers.append(fleet.imo(cells[draft.mapping["imo_number"]]))
        except (ValueError, IndexError):
            pass
    db.query(Vessel).filter(
        func.upper(func.trim(Vessel.imo_number)).in_(
            numbers + ["IMO " + n for n in numbers]
        )
    ).order_by(Vessel.id).with_for_update().all()
    findings, parsed = parse_rows(db, company.id, draft, draft.mapping)
    for finding, row in zip(findings, parsed):
        if row is not None:
            claims = fleet.declare(db, company, actor, row, import_id=draft.id)
            finding["status"] = (
                "CONFLICT"
                if any(c.status == "SUBMITTED" for c in claims)
                else "ACCEPTED"
            )
            finding["message"] = (
                "Accepted for aggregate data."
                if finding["status"] == "ACCEPTED"
                else "Conflicting roles are pending Super Admin review."
            )
            finding["entry_id"] = claims[0].id
    result = preview_result(draft, findings)
    result["source_sha256"] = draft.source_sha256
    draft.result = result
    draft.status = "IMPORTED"
    draft.imported_at = datetime.now(timezone.utc)
    fleet.audit(
        db,
        actor,
        company.id,
        draft.id,
        "fleet_import_completed",
        target_type="fleet_import",
        after={"counts": result["counts"], "source_sha256": draft.source_sha256},
    )
    db.commit()
    return result


@router.get("/imports/{import_id}/errors.csv")
def errors(
    import_id: int,
    db: Session = Depends(get_db),
    context=Depends(current_company_member),
):
    company = write_context(context)
    draft = owned_import(db, company.id, import_id)
    if not draft.mapping or draft.status == "UPLOADED":
        raise HTTPException(409, "Map and review the file first.")
    findings = (
        draft.result["findings"]
        if draft.result
        else parse_rows(db, company.id, draft, draft.mapping)[0]
    )
    rows = [LABELS + ["Source Row", "Status", "Error"]]
    for finding in findings:
        if finding["status"] in ("ACCEPTED", "CONFLICT"):
            continue
        cells = draft.rows[finding["row"] - 2]
        rows.append(
            [
                cells[draft.mapping[f]] if draft.mapping[f] < len(cells) else ""
                for f in fleet.FIELDS
            ]
            + [finding["row"], finding["status"], finding["message"]]
        )
    return Response(
        safe_csv(rows),
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="fleet-row-errors.csv"'},
    )
