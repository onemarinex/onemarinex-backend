"""Fleet registry and bounded file preview. Access still requires independent approval."""

import csv
import io
from datetime import datetime, timezone, date
from typing import Optional
from zipfile import ZipFile
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from pydantic import ValidationError
from sqlalchemy import or_
from sqlalchemy.orm import Session
from app.db.session import get_db
from app.db.models.vessel import Vessel
from app.db.models.shipping_company import VesselRoleClaim
from app.api.v1.routes_auth import get_current_user
from app.api.v1.routes_shipping_company import (
    current_company_member,
    _require_company_admin,
    preview_fleet,
    FleetPreviewIn,
    FleetRowIn,
)

router = APIRouter()


@router.get("/registry")
def registry(
    search: str = Query("", max_length=120),
    status: Optional[str] = None,
    offset: int = Query(0, ge=0),
    limit: int = Query(25, ge=1, le=100),
    db: Session = Depends(get_db),
    context=Depends(current_company_member),
):
    company, _ = context
    query = (
        db.query(VesselRoleClaim, Vessel)
        .join(Vessel, Vessel.id == VesselRoleClaim.vessel_id)
        .filter(VesselRoleClaim.company_id == company.id)
    )
    if search.strip():
        term = "%" + search.strip().replace("%", "\\%").replace("_", "\\_") + "%"
        query = query.filter(
            or_(
                Vessel.name.ilike(term, escape="\\"),
                Vessel.imo_number.ilike(term, escape="\\"),
            )
        )
    if status:
        if status not in {
            "APPROVED",
            "SUBMITTED",
            "UNDER_REVIEW",
            "MORE_EVIDENCE",
            "REJECTED",
            "REVOKED",
            "CORRECTED",
        }:
            raise HTTPException(422, "Choose a known review status.")
        query = query.filter(VesselRoleClaim.status == status)
    return {
        "total": query.count(),
        "offset": offset,
        "limit": limit,
        "items": [
            {
                "id": c.id,
                "vessel_id": c.vessel_id,
                "vessel_name": vessel.name,
                "imo_number": vessel.imo_number,
                "role": c.role,
                "effective_start": c.effective_start,
                "effective_end": c.effective_end,
                "status": c.status,
                "review_note": c.review_reason,
            }
            for c, vessel in query.order_by(Vessel.name, VesselRoleClaim.id)
            .offset(offset)
            .limit(limit)
        ],
    }


@router.post("/import-preview")
def import_preview(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    context=Depends(current_company_member),
    actor=Depends(get_current_user),
):
    _require_company_admin(context[1])
    raw = file.file.read(2_000_001)
    if len(raw) > 2_000_000:
        raise HTTPException(413, "Use a file of at most 2 MB.")
    filename = file.filename or ""
    try:
        if filename.lower().endswith(".csv"):
            data = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
        elif filename.lower().endswith(".xlsx"):
            from openpyxl import load_workbook

            with ZipFile(io.BytesIO(raw)) as archive:
                if (
                    len(archive.infolist()) > 1000
                    or sum(i.file_size for i in archive.infolist()) > 8_000_000
                ):
                    raise HTTPException(
                        413, "The workbook expands beyond the allowed size."
                    )
            workbook = load_workbook(
                io.BytesIO(raw), read_only=True, data_only=False, keep_links=False
            )
            try:
                sheet = workbook.active
                if sheet.max_row and sheet.max_row > 501:
                    raise HTTPException(422, "Use at most 500 data rows.")
                data = [
                    list(r)
                    for r in sheet.iter_rows(max_row=502, max_col=5, values_only=True)
                ]
            finally:
                workbook.close()
        else:
            raise HTTPException(422, "Use a CSV or XLSX file.")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(
            422, "The file could not be read. Use the supplied column template."
        ) from None
    data = [r for r in data if any(v is not None and str(v).strip() for v in r)]
    if not data or [str(v or "").strip().lower() for v in data[0][:4]] != [
        "imo",
        "role",
        "start",
        "end",
    ]:
        raise HTTPException(422, "The first row must be: IMO, role, start, end.")
    if len(data) > 501:
        raise HTTPException(422, "Use at most 500 data rows.")
    rows, valid, indexes, findings = [], [], [], []

    def stamp(value):
        if value in (None, ""):
            return None
        if isinstance(value, datetime):
            return value.replace(tzinfo=value.tzinfo or timezone.utc)
        if isinstance(value, date):
            return datetime.combine(value, datetime.min.time(), tzinfo=timezone.utc)
        value = str(value).strip()
        return (
            datetime.fromisoformat(value.replace("Z", "+00:00")).replace(
                tzinfo=timezone.utc
            )
            if len(value) == 10
            else datetime.fromisoformat(value.replace("Z", "+00:00"))
        )

    for n, cells in enumerate(data[1:], 1):
        cells = (cells + [None] * 4)[:4]
        imo = str(cells[0] or "").strip()
        try:
            if any(isinstance(v, str) and v.startswith("=") for v in cells):
                raise ValueError("Formula cells are not allowed")
            row = FleetRowIn(
                imo_number=imo,
                role=str(cells[1] or "").strip(),
                effective_start=stamp(cells[2]),
                effective_end=stamp(cells[3]),
            )
            rows.append(row.model_dump(mode="json"))
            valid.append(row)
            indexes.append(n)
        except (ValueError, TypeError, ValidationError):
            rows.append(None)
            findings.append(
                {
                    "row": n,
                    "imo_number": imo,
                    "ok": False,
                    "problem": "Invalid date, role value or formula cell.",
                    "vessel_id": None,
                }
            )
    if valid:
        result = preview_fleet(
            FleetPreviewIn(rows=valid, filename=filename[:255]), db, context, actor
        )
        findings.extend(
            dict(f.model_dump(), row=indexes[f.row - 1]) for f in result.findings
        )
    return {
        "rows": rows,
        "findings": sorted(findings, key=lambda f: f["row"]),
        "filename": filename[:255],
        "accepted": sum(f["ok"] for f in findings),
        "total": len(rows),
    }


# Changes are requests. Neither submission nor import can approve a claim.
from typing import Literal
from pydantic import BaseModel, Field
from app.api.v1.routes_company_onboarding import admin
from app.db.models.company_fleet_change import CompanyFleetChange
from app.db.models.shipping_company import CompanyAuditEvent, ShippingCompany
from app.services import company_claims as claims

admin_router = APIRouter()


class ChangeIn(BaseModel):
    model_config = {"str_strip_whitespace": True}
    action: Literal["EDIT", "END", "VOID"]
    effective_start: Optional[datetime] = None
    effective_end: Optional[datetime] = None
    reason: str = Field(min_length=10, max_length=4000)


class ChangeDecisionIn(BaseModel):
    model_config = {"str_strip_whitespace": True}
    approve: bool
    reason: str = Field(min_length=10, max_length=4000)


def change_out(db, change):
    claim = db.get(VesselRoleClaim, change.claim_id)
    vessel = db.get(Vessel, claim.vessel_id)
    return {
        "id": change.id,
        "company_name": db.get(ShippingCompany, change.company_id).legal_name,
        "claim_id": claim.id,
        "vessel_name": vessel.name,
        "imo_number": claim.imo_number,
        "role": claim.role,
        "action": change.action,
        "payload": change.payload,
        "reason": change.reason,
        "status": change.status,
        "decision_reason": change.decision_reason,
        "created_at": change.created_at,
    }


@router.post("/claims/{claim_id}/changes", status_code=201)
def request_change(
    claim_id: int,
    body: ChangeIn,
    db: Session = Depends(get_db),
    context=Depends(current_company_member),
    actor=Depends(get_current_user),
):
    _require_company_admin(context[1])
    claim = (
        db.query(VesselRoleClaim)
        .filter_by(id=claim_id, company_id=context[0].id)
        .with_for_update()
        .first()
    )
    if claim is None:
        raise HTTPException(404, "Fleet record not found.")
    if claim.status != "APPROVED" or claim.revoked_at:
        raise HTTPException(409, "Only an approved fleet record can be changed.")
    if (
        db.query(CompanyFleetChange.id)
        .filter_by(claim_id=claim.id, status="REQUESTED")
        .first()
    ):
        raise HTTPException(409, "A change is already awaiting review for this record.")
    if len(body.reason.strip()) < 10:
        raise HTTPException(422, "Give a reason of at least ten characters.")
    start = body.effective_start or claim.effective_start
    end = body.effective_end or claim.effective_end
    if body.action in ("EDIT", "END"):
        if start.tzinfo is None or (
            end is not None and (end.tzinfo is None or end <= start)
        ):
            raise HTTPException(422, "Use a valid effective period with time zones.")
        if body.action == "END" and body.effective_end is None:
            raise HTTPException(422, "Enter the effective end date.")
        if body.action == "END":
            start = claim.effective_start
    change = CompanyFleetChange(
        company_id=context[0].id,
        claim_id=claim.id,
        requested_by_user_id=actor.id,
        action=body.action,
        payload={
            "effective_start": start.isoformat(),
            "effective_end": end.isoformat() if end else None,
        },
        reason=body.reason.strip(),
    )
    db.add(change)
    db.flush()
    db.add(
        CompanyAuditEvent(
            company_id=context[0].id,
            actor_user_id=actor.id,
            target_type="fleet_change",
            target_id=change.id,
            action="fleet_change_requested",
            after_state=change.payload,
            reason=change.reason,
        )
    )
    db.commit()
    return change_out(db, change)


@router.get("/changes")
def company_changes(
    offset: int = Query(0, ge=0),
    limit: int = Query(25, ge=1, le=100),
    db: Session = Depends(get_db),
    context=Depends(current_company_member),
):
    q = db.query(CompanyFleetChange).filter_by(company_id=context[0].id)
    return {
        "total": q.count(),
        "items": [
            change_out(db, r)
            for r in q.order_by(CompanyFleetChange.id.desc())
            .offset(offset)
            .limit(limit)
        ],
    }


@admin_router.get("/fleet-changes")
def fleet_changes(
    offset: int = Query(0, ge=0),
    limit: int = Query(25, ge=1, le=100),
    db: Session = Depends(get_db),
    actor=Depends(admin),
):
    q = db.query(CompanyFleetChange)
    return {
        "total": q.count(),
        "items": [
            change_out(db, r)
            for r in q.order_by(CompanyFleetChange.id.desc())
            .offset(offset)
            .limit(limit)
        ],
    }


@admin_router.post("/fleet-changes/{change_id}/decision")
def decide_change(
    change_id: int,
    body: ChangeDecisionIn,
    db: Session = Depends(get_db),
    actor=Depends(admin),
):
    change = (
        db.query(CompanyFleetChange).filter_by(id=change_id).with_for_update().first()
    )
    if change is None:
        raise HTTPException(404, "Change request not found.")
    if change.status != "REQUESTED":
        raise HTTPException(409, "This change was already decided.")
    claim = (
        db.query(VesselRoleClaim).filter_by(id=change.claim_id).with_for_update().one()
    )
    if body.approve:
        if claim.status != "APPROVED" or claim.revoked_at:
            raise HTTPException(
                409, "The source fleet record has changed. Review a new request."
            )
        try:
            if change.action == "VOID":
                claims.revoke_claim(db, claim=claim, actor=actor, reason=body.reason)
            else:
                replacement = claims.correct_claim(
                    db,
                    claim=claim,
                    actor=actor,
                    reason=body.reason,
                    effective_start=datetime.fromisoformat(
                        change.payload["effective_start"]
                    ),
                    effective_end=(
                        datetime.fromisoformat(change.payload["effective_end"])
                        if change.payload["effective_end"]
                        else None
                    ),
                )
                change.replacement_claim_id = replacement.id
        except claims.ClaimError as exc:
            db.rollback()
            raise HTTPException(409, str(exc)) from None
    change.status = "APPROVED" if body.approve else "REJECTED"
    change.decided_by_user_id = actor.id
    change.decision_reason = body.reason.strip()
    change.decided_at = datetime.now(timezone.utc)
    db.add(
        CompanyAuditEvent(
            company_id=change.company_id,
            actor_user_id=actor.id,
            target_type="fleet_change",
            target_id=change.id,
            action="fleet_change_" + change.status.lower(),
            after_state=change.payload,
            reason=change.decision_reason,
        )
    )
    db.commit()
    return change_out(db, change)
